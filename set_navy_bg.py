from docx import Document
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

doc = Document(r"C:\Users\nehas5\OneDrive - AMDOCS\Backup Folders\Desktop\DAM\Resume_Vikram_Hardas.docx")

bg_xml = '<w:background %s w:color="1B2A4A"/>' % nsdecls("w")
bg = parse_xml(bg_xml)
doc.element.insert(0, bg)

settings = doc.settings.element
dpbg_xml = '<w:displayBackgroundShape %s/>' % nsdecls("w")
dpbg = parse_xml(dpbg_xml)
settings.append(dpbg)

# Make all text white so it's visible on navy background
from docx.shared import RGBColor
for para in doc.paragraphs:
    for run in para.runs:
        if run.font.color.rgb is None or run.font.color.rgb == RGBColor(0, 0, 0):
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    if run.font.color.rgb is None or run.font.color.rgb == RGBColor(0, 0, 0):
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

output = r"C:\Users\nehas5\OneDrive - AMDOCS\Backup Folders\Desktop\DAM\Resume_Vikram_Hardas_NavyBlue.docx"
doc.save(output)
print(f"Saved: {output}")
